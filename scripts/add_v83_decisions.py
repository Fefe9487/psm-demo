"""Preserve the V8.3 draft and append project-specific initial filing decisions."""

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "石化及化學工廠製程安全管理參考手冊-V8.3.docx"
OUTPUT = ROOT / "docs" / "V8.3手冊_初次填報系統決議補充版.docx"


def set_font(run, *, bold=False, size=10.5):
    run.bold = bold
    run.font.name = "Microsoft JhengHei"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    fonts = run._element.get_or_add_rPr().rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run._element.get_or_add_rPr().insert(0, fonts)
    fonts.set(qn("w:eastAsia"), "Microsoft JhengHei")


def paragraph(doc, label, body):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.25
    set_font(p.add_run(label), bold=True)
    set_font(p.add_run(body))
    return p


doc = Document(SOURCE)

# This is a project decision, so leave the source draft untouched and point to
# the clearly labeled appendix in this deliverable copy.
cover_note = doc.paragraphs[3].insert_paragraph_before()
cover_note.paragraph_format.space_before = Pt(8)
cover_note.paragraph_format.space_after = Pt(8)
set_font(
    cover_note.add_run("本補充版文末附有初次填報系統決議；原 V8.3 草案內容保留。"),
    size=10,
)

doc.add_page_break()
heading = doc.add_paragraph("附錄二 初次填報系統決議", style="Heading 1")
heading.paragraph_format.keep_with_next = True
for run in heading.runs:
    set_font(run, bold=True, size=15)

paragraph(
    doc,
    "適用範圍　",
    "決議日期為 2026 年 9 月 27 日。本附錄記錄本專案填報與驗收規則，"
    "不改寫前文草案，也不代替現行法令或主管機關審查要求。",
)
paragraph(
    doc,
    "產品階段與範圍　",
    "現階段系統只支援初次填報 PSM 十四項；甲類申請書與安全衛生管理基本資料不在本系統送件範圍。"
    "日後製程修改或定期審查，事業單位應以本系統既有資料為起點，重新檢視、更新並填寫當次資料。"
    "本階段不把後續填報視為初次申請，也不自動沿用舊資料作為已確認內容。",
)
paragraph(
    doc,
    "1.5 設備安全上下限　",
    "所有設備均須逐台檢討操作與安全上下限。若一台設備確實沒有適用的壓力、溫度、流量或組成檢討項目，"
    "可選擇全部不適用，但必須填寫該設備的具體理由；系統不得以空白項目自動視為不適用。",
)
paragraph(
    doc,
    "7.1 啟動前安全檢查　",
    "所有設備均須逐台確認本次是否適用啟動前安全檢查。適用者填檢查與改善追蹤紀錄；"
    "不適用者填逐台理由。全場所確無適用案件時，另填全場所無案件聲明與理由；"
    "該聲明不可取代逐台判定，也不得與適用設備或檢查紀錄並存。",
)
paragraph(
    doc,
    "8.2 安全關鍵設備失效率　",
    "逐台區分有效數值、運行時間為零、資料不足三種情況。運行時間為零或資料不足均須註明原因，"
    "不計算該台失效率，且不納入同類設備失效率分子與分母。資料不足可以保存並顯示於報表，"
    "但章節須標示待人工確認，不得自動顯示基本檢核通過；不得用臆測數值補足缺漏。",
)
paragraph(
    doc,
    "12.1 緊急應變演練　",
    "每一個已列緊急狀況都須對應演練計畫的涵蓋項目。已執行演練者記錄日期並附紀錄；"
    "初次填報時尚未執行者，填預定日期與未執行原因。尚未執行的情境可保存，"
    "但標示待人工確認，不得以空白演練欄位自動視為完成。",
)
paragraph(
    doc,
    "第六章承攬商　",
    "即使近五年沒有承攬事故，仍須登錄所有承攬商與承攬期間，並另行據實聲明零事故。",
)
paragraph(
    doc,
    "第十章 MOC 文件　",
    "初次填報所稱「內部管理程序」與「變更管理計畫書」指同一份管理 MOC 運作的程序書；"
    "事業單位上傳或引用該程序書即可，不因名稱不同要求重複上傳。",
)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
